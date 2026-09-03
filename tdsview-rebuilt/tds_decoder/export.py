from __future__ import annotations

import json
from collections import Counter
from datetime import datetime
from html import escape
from pathlib import Path
from typing import Callable, Iterable, TextIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .decoder import environment_rows
from .models import DecodedDataset, EventRecord


EVENT_HEADERS = [
    "Ref Nr",
    "Start time",
    "End Time",
    "Duration",
    "Dist Text",
    "ECode 0",
    "ECode 0 Descr",
    "Event Name",
    "EnvBl Id",
    "Cnt",
    "Ev Id",
    "Limit",
    "Vh Pos",
    "Prc Id",
    "Subsys",
    "Subsys Desc",
    "Prio",
    "ECode 1",
    "ECode 1 Descr",
    "ECode 2",
    "ECode 2 Descr",
    "ECode 3",
    "ECode 3 Descr",
    "UniqueRef",
    "Vehicle Name",
    "User1",
    "User2",
    "User3",
    "User4",
    "User5",
    "Position",
    "Altitude",
    "Speed",
    "Heading",
    "UTC Time",
    "Odometer",
    "Trip",
    "rep_prio_active",
    "rep_prio_passive",
    "extra_prio",
    "extra attrib1",
    "extra attrib2",
    "extra attrib3",
    "extra attrib4",
    "extra attrib5",
]

HEADER_FILL = PatternFill("solid", fgColor="FFF200")
CHANGED_FILL = PatternFill("solid", fgColor="FFD700")
INFERRED_FILL = PatternFill("solid", fgColor="FFE6B3")
HEADER_FONT = Font(bold=True)


def _html(value: object) -> str:
    return escape("" if value is None else str(value), quote=True)


def _time(value) -> str:
    return value.strftime("%Y-%m-%d %H:%M:%S") if value else ""


def _duration(record: EventRecord) -> str:
    if record.end_time is None:
        return ""
    delta = record.end_time - record.start_time
    if delta.total_seconds() < 0:
        return "Negative!"
    days = delta.days
    seconds = delta.seconds
    prefix = f"{days}." if days else ""
    return f"{prefix}{seconds // 3600:02d}:{(seconds % 3600) // 60:02d}:{seconds % 60:02d}"


def _compact_float(value: float) -> str:
    if value == int(value):
        return str(int(value))
    return format(value, ".7g")


def event_export_row(dataset: DecodedDataset, record: EventRecord, source_index: int = 0) -> list[str]:
    definition = record.definition
    subsystem = dataset.definitions.subsystems.get(record.subsystem_number)
    descriptions = [code.description if code else "" for code in record.code_definitions]
    extra = definition.extra_attributes if definition else ("", "", "", "", "")
    return [
        f"{record.reference_number:05d}",
        _time(record.start_time),
        _time(record.end_time),
        _duration(record),
        definition.description if definition else "",
        f"{record.error_codes[0]:04X}",
        descriptions[0],
        definition.name if definition else "",
        definition.environment_block_id if definition else "",
        f"{record.event_count:05d}",
        f"{record.event_id:05d}",
        str(record.limit),
        str(record.vehicle_position),
        f"{record.process_id:03d}",
        subsystem.name if subsystem else str(record.subsystem_number),
        subsystem.description if subsystem else str(record.subsystem_number),
        str(record.priority),
        f"{record.error_codes[1]:04X}",
        descriptions[1],
        f"{record.error_codes[2]:04X}",
        descriptions[2],
        f"{record.error_codes[3]:04X}",
        descriptions[3],
        f"{source_index}:{record.unique_reference}",
        dataset.edv.header.vehicle_name,
        "",
        "",
        "",
        "",
        "",
        f"{record.latitude:.6f}/{record.longitude:.6f}",
        _compact_float(record.altitude),
        _compact_float(record.speed),
        _compact_float(record.heading),
        str(record.utc_time),
        str(record.odometer),
        str(record.trip),
        definition.rep_prio_active if definition else "",
        definition.rep_prio_passive if definition else "",
        definition.extra_prio if definition else "",
        *extra,
    ]


def file_information_rows(dataset: DecodedDataset) -> list[tuple[str, object]]:
    header = dataset.edv.header
    return [
        ("ED_V file", str(dataset.edv.source_path)),
        ("File ID", header.file_id),
        ("ED protocol", header.ed_version),
        ("Project", header.project_name),
        ("ED_V project version", header.project_version),
        ("ED_D file", str(dataset.definitions.source_path)),
        ("ED_D project version", dataset.definitions.project_version),
        ("OTI schema version", dataset.definitions.schema_version),
        ("Vehicle", header.vehicle_name),
        ("ODBS address", header.odbs_address),
        ("Readout date", header.readout_date),
        ("Readout time", header.readout_time),
        ("Event records", len(dataset.edv.records)),
        ("Mapped event definitions", dataset.mapped_event_count),
        ("Decodable environments", dataset.environment_event_count),
        ("Four-byte trailer", dataset.edv.trailer.hex().upper()),
    ]


def _header_row(sheet, headers: Iterable[str]):
    cells = []
    for value in headers:
        cell = WriteOnlyCell(sheet, value=value)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cells.append(cell)
    sheet.append(cells)


def _set_widths(sheet, widths: list[float]):
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _attributes(values: dict[str, object]) -> str:
    rendered = []
    for key, value in values.items():
        if value is None or value == "":
            continue
        rendered.append(f' {key}="{_html(value)}"')
    return "".join(rendered)


def _write_table_start(handle: TextIO, table_id: str, headers: Iterable[str]):
    handle.write(f'<div class="table-wrap"><table id="{_html(table_id)}">\n<thead><tr>')
    for header in headers:
        handle.write(f"<th>{_html(header)}</th>")
    handle.write("</tr></thead>\n<tbody>\n")


def _write_table_end(handle: TextIO):
    handle.write("</tbody>\n</table></div>\n")


def _write_cells(
    handle: TextIO,
    values: Iterable[object],
    *,
    tag: str = "td",
    classes: Iterable[str] | None = None,
):
    class_values = list(classes or [])
    for index, value in enumerate(values):
        class_name = class_values[index] if index < len(class_values) else ""
        class_attr = f' class="{_html(class_name)}"' if class_name else ""
        handle.write(f"<{tag}{class_attr}>{_html(value)}</{tag}>")


def _write_report_head(handle: TextIO, title: str):
    handle.write("<!doctype html>\n<html lang=\"en\">\n<head>\n")
    handle.write("<meta charset=\"utf-8\">\n")
    handle.write("<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">\n")
    handle.write(f"<title>{_html(title)}</title>\n")
    handle.write(
        """
<style>
:root {
  color-scheme: light;
  --blue: #0b5cab;
  --blue-dark: #063b73;
  --blue-soft: #eaf3ff;
  --line: #9cbfe5;
  --text: #17324d;
  --muted: #536a80;
  --changed: #ffe88a;
  --inferred: #fff5d7;
  --warning: #f8df8f;
}
* {
  box-sizing: border-box;
}
body {
  margin: 0;
  background: #f7f9fc;
  color: var(--text);
  font-family: "Segoe UI", Arial, sans-serif;
  font-size: 13px;
  line-height: 1.45;
}
header {
  background: #ffffff;
  border-bottom: 1px solid var(--line);
  padding: 18px 24px 14px;
}
h1 {
  margin: 0 0 6px;
  color: var(--blue-dark);
  font-size: 24px;
  letter-spacing: 0;
}
h2 {
  margin: 28px 0 10px;
  color: var(--blue-dark);
  font-size: 18px;
  letter-spacing: 0;
}
p {
  margin: 0 0 8px;
}
nav {
  position: sticky;
  top: 0;
  z-index: 3;
  display: flex;
  flex-wrap: wrap;
  gap: 8px;
  padding: 10px 24px;
  background: #ffffff;
  border-bottom: 1px solid var(--line);
}
nav a {
  color: #ffffff;
  background: var(--blue);
  padding: 6px 10px;
  border-radius: 4px;
  text-decoration: none;
  font-weight: 600;
}
main {
  padding: 0 24px 32px;
}
.muted {
  color: var(--muted);
}
.stats {
  display: grid;
  grid-template-columns: repeat(auto-fit, minmax(180px, 1fr));
  gap: 8px;
  margin: 16px 0 4px;
}
.stat {
  background: #ffffff;
  border: 1px solid var(--line);
  border-radius: 4px;
  padding: 10px 12px;
}
.stat strong {
  display: block;
  color: var(--blue-dark);
  font-size: 20px;
}
.notice {
  margin: 10px 0;
  padding: 10px 12px;
  background: var(--warning);
  border: 1px solid #c49b2b;
  border-radius: 4px;
}
.table-wrap {
  max-height: 76vh;
  overflow: auto;
  background: #ffffff;
  border: 1px solid var(--line);
}
.table-controls {
  display: flex;
  flex-wrap: wrap;
  align-items: center;
  gap: 8px;
  margin: 0 0 8px;
}
.table-controls input,
.table-controls select,
.table-controls button {
  min-height: 32px;
  border: 1px solid var(--line);
  border-radius: 4px;
  background: #ffffff;
  color: var(--text);
  padding: 5px 8px;
  font: inherit;
}
.table-controls input {
  min-width: 280px;
}
.table-controls button {
  background: var(--blue);
  color: #ffffff;
  cursor: pointer;
  font-weight: 600;
}
.table-controls button:disabled {
  cursor: default;
  opacity: 0.5;
}
.page-status {
  color: var(--muted);
  font-weight: 600;
}
.report-error {
  margin: 12px 24px;
  padding: 12px;
  border: 1px solid #a11;
  border-radius: 4px;
  background: #ffe8e8;
  color: #710000;
}
[hidden] {
  display: none !important;
}
table {
  width: max-content;
  min-width: 100%;
  border-collapse: collapse;
}
th,
td {
  border: 1px solid var(--line);
  padding: 5px 7px;
  vertical-align: top;
  white-space: nowrap;
}
th {
  position: sticky;
  top: 0;
  z-index: 2;
  background: var(--blue);
  color: #ffffff;
  text-align: left;
}
tbody tr:nth-child(even) td {
  background: var(--blue-soft);
}
tbody tr.inferred td {
  background: var(--inferred);
}
td.changed,
tbody tr.inferred td.changed {
  background: var(--changed);
  font-weight: 700;
}
td.pre {
  white-space: pre-wrap;
  min-width: 420px;
}
footer {
  padding: 18px 24px;
  color: var(--muted);
  border-top: 1px solid var(--line);
  background: #ffffff;
}
@media print {
  nav {
    display: none;
  }
  .table-wrap {
    max-height: none;
    overflow: visible;
  }
  th {
    position: static;
  }
}
</style>
</head>
<body>
"""
    )


def _export_expanded_html(
    dataset: DecodedDataset,
    path: str | Path,
    progress: Callable[[str], None] | None = None,
) -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    header = dataset.edv.header
    generated = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    title = f"{dataset.edv.source_path.stem} - TDS HTML Report"
    max_samples = max(
        (record.environment_block_count_start for record in dataset.edv.records), default=0
    )

    with output.open("w", encoding="utf-8", newline="\n") as handle:
        _write_report_head(handle, title)
        handle.write("<header>\n")
        handle.write("<h1>BTIL TDS Event & Environment Report</h1>\n")
        handle.write(
            f"<p class=\"muted\">Generated { _html(generated) } from { _html(dataset.edv.source_path.name) }</p>\n"
        )
        handle.write("<div class=\"stats\">\n")
        for label, value in [
            ("Total events", len(dataset.edv.records)),
            ("Named events", dataset.mapped_event_count),
            ("Decoded environments", dataset.environment_event_count),
            ("Vehicle", header.vehicle_name),
            ("Project", header.project_name),
            ("ED_V version", header.project_version),
            ("ED_D version", dataset.definitions.project_version),
            ("ODBS", header.odbs_address),
        ]:
            handle.write("<div class=\"stat\">")
            handle.write(f"<span>{_html(label)}</span><strong>{_html(value)}</strong>")
            handle.write("</div>\n")
        handle.write("</div>\n")
        if dataset.warnings:
            handle.write("<div class=\"notice\"><strong>Warnings</strong><ul>")
            for warning in dataset.warnings:
                handle.write(f"<li>{_html(warning)}</li>")
            handle.write("</ul></div>\n")
        handle.write("</header>\n")

        handle.write(
            """
<nav aria-label="Report sections">
  <a href="#file-information">File Information</a>
  <a href="#event-list">Event List</a>
  <a href="#environment-data">Environment Data</a>
  <a href="#fault-summary">Fault Summary</a>
  <a href="#repair-text">Repair Text</a>
</nav>
<main>
"""
        )

        handle.write('<section id="file-information">\n<h2>File Information</h2>\n')
        _write_table_start(handle, "file-info-table", ["Field", "Value"])
        for key, value in file_information_rows(dataset):
            handle.write("<tr>")
            _write_cells(handle, [key, value])
            handle.write("</tr>\n")
        for warning in dataset.warnings:
            handle.write("<tr>")
            _write_cells(handle, ["Warning", warning])
            handle.write("</tr>\n")
        _write_table_end(handle)
        handle.write("</section>\n")

        if progress:
            progress("Writing HTML event list...")
        handle.write('<section id="event-list">\n<h2>Event List</h2>\n')
        _write_table_start(handle, "event-list-table", EVENT_HEADERS)
        for record in dataset.edv.records:
            attrs = _attributes(
                {
                    "data-reference": f"{record.reference_number:05d}",
                    "data-process": f"{record.process_id:02X}",
                    "data-event": f"{record.event_id:03X}",
                }
            )
            handle.write(f"<tr{attrs}>")
            _write_cells(handle, event_export_row(dataset, record))
            handle.write("</tr>\n")
        _write_table_end(handle)
        handle.write("</section>\n")

        if progress:
            progress("Writing HTML environment data...")
        environment_headers = [
            "Ref Nr",
            "Start Time",
            "Event Name",
            "Event Description",
            "Process",
            "Event Id",
            "Env Block",
            "Mapping",
            "Variable Name",
            "Description",
            "Unit",
        ]
        for sample in range(1, max_samples + 1):
            environment_headers.extend([f"Offset {sample} (ms)", f"Value {sample}"])
        handle.write('<section id="environment-data">\n<h2>Environment Data</h2>\n')
        _write_table_start(handle, "environment-data-table", environment_headers)
        for record_index, record in enumerate(dataset.edv.records, 1):
            event_name = (
                record.definition.name
                if record.definition
                else f"Unknown {record.process_id:02X}:{record.event_id:03X}"
            )
            event_description = record.definition.description if record.definition else ""
            block_name = record.environment_block.block_id if record.environment_block else ""
            row_class = "inferred" if record.environment_mapping != "exact" else ""
            attrs = _attributes(
                {
                    "class": row_class,
                    "data-reference": f"{record.reference_number:05d}",
                    "data-process": f"{record.process_id:02X}",
                    "data-event": f"{record.event_id:03X}",
                    "data-mapping": record.environment_mapping,
                }
            )
            for environment in environment_rows(record):
                values: list[object] = [
                    f"{record.reference_number:05d}",
                    _time(record.start_time),
                    event_name,
                    event_description,
                    f"{record.process_id:02X}",
                    f"{record.event_id:03X}",
                    block_name,
                    record.environment_mapping,
                    environment.name,
                    environment.description,
                    environment.unit,
                ]
                classes = [""] * len(values)
                for sample in range(max_samples):
                    if sample < len(environment.values):
                        values.extend([environment.offsets_ms[sample], environment.values[sample]])
                        classes.extend(["", "changed" if environment.changed[sample] else ""])
                    else:
                        values.extend(["", ""])
                        classes.extend(["", ""])
                handle.write(f"<tr{attrs}>")
                _write_cells(handle, values, classes=classes)
                handle.write("</tr>\n")
            if progress and record_index % 50 == 0:
                progress(f"HTML environment data: {record_index}/{len(dataset.edv.records)} events")
        _write_table_end(handle)
        handle.write("</section>\n")

        handle.write('<section id="fault-summary">\n<h2>Fault Summary</h2>\n')
        _write_table_start(handle, "fault-summary-table", ["Count", "Process", "Event Id", "Event Name", "Description"])
        counts = Counter((record.process_id, record.event_id) for record in dataset.edv.records)
        for (process_id, event_id), count in counts.most_common():
            definition = dataset.definitions.events.get((process_id, event_id))
            handle.write("<tr>")
            _write_cells(
                handle,
                [
                    count,
                    f"{process_id:02X}",
                    f"{event_id:03X}",
                    definition.name if definition else "",
                    definition.description if definition else "Definition unavailable",
                ],
            )
            handle.write("</tr>\n")
        _write_table_end(handle)
        handle.write("</section>\n")

        handle.write('<section id="repair-text">\n<h2>Repair Text</h2>\n')
        _write_table_start(handle, "repair-text-table", ["Event Name", "Event Description", "Repair / Cause / Remedy"])
        used_names: set[str] = set()
        for record in dataset.edv.records:
            if not record.definition or record.definition.name in used_names:
                continue
            used_names.add(record.definition.name)
            details = dataset.texts.details_for(record.definition.name)
            handle.write("<tr>")
            _write_cells(
                handle,
                [record.definition.name, record.definition.description, details],
                classes=["", "", "pre"],
            )
            handle.write("</tr>\n")
        _write_table_end(handle)
        handle.write("</section>\n")
        handle.write("</main>\n")
        handle.write(
            f"<footer>Generated by BTIL Data Analyser. Source ED_D: {_html(dataset.definitions.source_path)}</footer>\n"
        )
        handle.write("</body>\n</html>\n")

    if progress:
        progress(f"Saved {output}")
    return output


def _compact_html_payload(
    dataset: DecodedDataset,
    max_samples: int,
    progress: Callable[[str], None] | None = None,
) -> dict[str, object]:
    """Build a lossless, deduplicated payload for the standalone HTML report."""
    strings: list[str] = []
    string_indexes: dict[str, int] = {}

    def intern(value: object) -> int:
        text = "" if value is None else str(value)
        existing = string_indexes.get(text)
        if existing is not None:
            return existing
        index = len(strings)
        strings.append(text)
        string_indexes[text] = index
        return index

    def indexed(values: Iterable[object]) -> list[int]:
        return [intern(value) for value in values]

    file_rows = [indexed(row) for row in file_information_rows(dataset)]
    file_rows.extend(indexed(["Warning", warning]) for warning in dataset.warnings)

    if progress:
        progress("Indexing compact HTML event list...")
    event_rows = [indexed(event_export_row(dataset, record)) for record in dataset.edv.records]

    environment_headers = [
        "Ref Nr",
        "Start Time",
        "Event Name",
        "Event Description",
        "Process",
        "Event Id",
        "Env Block",
        "Mapping",
        "Variable Name",
        "Description",
        "Unit",
    ]
    for sample in range(1, max_samples + 1):
        environment_headers.extend([f"Offset {sample} (ms)", f"Value {sample}"])

    if progress:
        progress("Indexing compact HTML environment data...")
    environment_groups: list[list[object]] = []
    for record_index, record in enumerate(dataset.edv.records, 1):
        rows = list(environment_rows(record))
        if not rows:
            continue
        event_name = (
            record.definition.name
            if record.definition
            else f"Unknown {record.process_id:02X}:{record.event_id:03X}"
        )
        metadata = indexed(
            [
                f"{record.reference_number:05d}",
                _time(record.start_time),
                event_name,
                record.definition.description if record.definition else "",
                f"{record.process_id:02X}",
                f"{record.event_id:03X}",
                record.environment_block.block_id if record.environment_block else "",
                record.environment_mapping,
            ]
        )
        default_offsets = list(rows[0].offsets_ms)
        compact_rows: list[list[object]] = []
        for environment in rows:
            compact_row: list[object] = [
                intern(environment.name),
                intern(environment.description),
                intern(environment.unit),
                indexed(environment.values),
                [index for index, changed in enumerate(environment.changed) if changed],
            ]
            offsets = list(environment.offsets_ms)
            if offsets != default_offsets:
                compact_row.append(offsets)
            compact_rows.append(compact_row)
        environment_groups.append([metadata, default_offsets, compact_rows])
        if progress and record_index % 50 == 0:
            progress(f"Compact HTML data: {record_index}/{len(dataset.edv.records)} events")

    counts = Counter((record.process_id, record.event_id) for record in dataset.edv.records)
    summary_rows = []
    for (process_id, event_id), count in counts.most_common():
        definition = dataset.definitions.events.get((process_id, event_id))
        summary_rows.append(
            indexed(
                [
                    count,
                    f"{process_id:02X}",
                    f"{event_id:03X}",
                    definition.name if definition else "",
                    definition.description if definition else "Definition unavailable",
                ]
            )
        )

    repair_rows = []
    used_names: set[str] = set()
    for record in dataset.edv.records:
        if not record.definition or record.definition.name in used_names:
            continue
        used_names.add(record.definition.name)
        repair_rows.append(
            indexed(
                [
                    record.definition.name,
                    record.definition.description,
                    dataset.texts.details_for(record.definition.name),
                ]
            )
        )

    return {
        "v": 1,
        "s": strings,
        "fh": indexed(["Field", "Value"]),
        "fi": file_rows,
        "eh": indexed(EVENT_HEADERS),
        "e": event_rows,
        "hh": indexed(environment_headers),
        "g": environment_groups,
        "m": max_samples,
        "sh": indexed(["Count", "Process", "Event Id", "Event Name", "Description"]),
        "su": summary_rows,
        "rh": indexed(["Event Name", "Event Description", "Repair / Cause / Remedy"]),
        "r": repair_rows,
    }


def _safe_json_for_html(value: object) -> str:
    # Script elements are raw-text containers. Escaping these characters keeps
    # OTI/user text from terminating the JSON element or becoming HTML markup.
    return (
        json.dumps(value, ensure_ascii=False, separators=(",", ":"))
        .replace("&", "\\u0026")
        .replace("<", "\\u003c")
        .replace(">", "\\u003e")
    )


def _write_dynamic_table_shell(handle: TextIO, table_id: str):
    handle.write(
        f'<div class="table-wrap"><table id="{_html(table_id)}">'
        "<thead></thead><tbody></tbody></table></div>\n"
    )


def export_html(
    dataset: DecodedDataset,
    path: str | Path,
    progress: Callable[[str], None] | None = None,
) -> Path:
    """Write a compact, complete, dependency-free interactive HTML report."""
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    header = dataset.edv.header
    generated = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    title = f"{dataset.edv.source_path.stem} - TDS HTML Report"
    max_samples = max(
        (record.environment_block_count_start for record in dataset.edv.records), default=0
    )
    payload = _compact_html_payload(dataset, max_samples, progress)
    payload_json = _safe_json_for_html(payload)

    with output.open("w", encoding="utf-8", newline="\n") as handle:
        _write_report_head(handle, title)
        handle.write("<header>\n")
        handle.write("<h1>BTIL TDS Event & Environment Report</h1>\n")
        handle.write(
            f'<p class="muted">Generated {_html(generated)} from '
            f"{_html(dataset.edv.source_path.name)}</p>\n"
        )
        handle.write('<div class="stats">\n')
        for label, value in [
            ("Total events", len(dataset.edv.records)),
            ("Named events", dataset.mapped_event_count),
            ("Decoded environments", dataset.environment_event_count),
            ("Vehicle", header.vehicle_name),
            ("Project", header.project_name),
            ("ED_V version", header.project_version),
            ("ED_D version", dataset.definitions.project_version),
            ("ODBS", header.odbs_address),
        ]:
            handle.write(
                f'<div class="stat"><span>{_html(label)}</span>'
                f"<strong>{_html(value)}</strong></div>\n"
            )
        handle.write("</div>\n")
        if dataset.warnings:
            handle.write('<div class="notice"><strong>Warnings</strong><ul>')
            for warning in dataset.warnings:
                handle.write(f"<li>{_html(warning)}</li>")
            handle.write("</ul></div>\n")
        handle.write("</header>\n")
        handle.write(
            """
<nav aria-label="Report sections">
  <a href="#file-information">File Information</a>
  <a href="#event-list">Event List</a>
  <a href="#environment-data">Environment Data</a>
  <a href="#fault-summary">Fault Summary</a>
  <a href="#repair-text">Repair Text</a>
</nav>
<div id="report-error" class="report-error" hidden></div>
<noscript><div class="report-error">JavaScript is required to display this standalone report.</div></noscript>
<main>
<section id="file-information"><h2>File Information</h2>
"""
        )
        _write_dynamic_table_shell(handle, "file-info-table")
        handle.write('</section>\n<section id="event-list"><h2>Event List</h2>\n')
        _write_dynamic_table_shell(handle, "event-list-table")
        handle.write(
            """</section>
<section id="environment-data"><h2>Environment Data</h2>
<div class="table-controls">
  <label for="environment-search">Search all data</label>
  <input id="environment-search" type="search" placeholder="Reference, event, variable, or value">
  <button id="environment-search-button" type="button">Search</button>
  <button id="environment-clear-button" type="button">Clear</button>
  <label for="environment-page-size">Rows per page</label>
  <select id="environment-page-size">
    <option value="100">100</option>
    <option value="250" selected>250</option>
    <option value="500">500</option>
    <option value="1000">1000</option>
    <option value="0">Show all</option>
  </select>
  <button id="environment-previous" type="button">Previous</button>
  <button id="environment-next" type="button">Next</button>
  <span id="environment-page-info" class="page-status"></span>
</div>
"""
        )
        _write_dynamic_table_shell(handle, "environment-data-table")
        handle.write('</section>\n<section id="fault-summary"><h2>Fault Summary</h2>\n')
        _write_dynamic_table_shell(handle, "fault-summary-table")
        handle.write('</section>\n<section id="repair-text"><h2>Repair Text</h2>\n')
        _write_dynamic_table_shell(handle, "repair-text-table")
        handle.write("</section>\n</main>\n")
        handle.write(
            f"<footer>Generated by BTIL Data Analyser. Source ED_D: "
            f"{_html(dataset.definitions.source_path)}</footer>\n"
        )
        handle.write(
            f'<script id="report-data" type="application/json">{payload_json}</script>\n'
        )
        handle.write(
            r"""
<script>
(() => {
  "use strict";
  try {
    const report = JSON.parse(document.getElementById("report-data").textContent);
    const stringValue = index => index === null || index === undefined ? "" : (report.s[index] ?? "");

    function renderHeader(tableId, headerIndexes) {
      const head = document.querySelector(`#${tableId} thead`);
      const row = document.createElement("tr");
      for (const index of headerIndexes) {
        const cell = document.createElement("th");
        cell.textContent = stringValue(index);
        row.appendChild(cell);
      }
      head.replaceChildren(row);
    }

    function appendRow(body, values, classes = [], rowClass = "") {
      const row = document.createElement("tr");
      if (rowClass) row.className = rowClass;
      values.forEach((value, index) => {
        const cell = document.createElement("td");
        cell.textContent = value === null || value === undefined ? "" : String(value);
        if (classes[index]) cell.className = classes[index];
        row.appendChild(cell);
      });
      body.appendChild(row);
      return row;
    }

    function renderIndexedTable(tableId, headerIndexes, rows, preColumn = -1) {
      renderHeader(tableId, headerIndexes);
      const body = document.querySelector(`#${tableId} tbody`);
      const fragment = document.createDocumentFragment();
      for (const row of rows) {
        const classes = Array(row.length).fill("");
        if (preColumn >= 0) classes[preColumn] = "pre";
        appendRow(fragment, row.map(stringValue), classes);
      }
      body.replaceChildren(fragment);
    }

    renderIndexedTable("file-info-table", report.fh, report.fi);
    renderIndexedTable("event-list-table", report.eh, report.e);
    renderIndexedTable("fault-summary-table", report.sh, report.su);
    renderIndexedTable("repair-text-table", report.rh, report.r, 2);
    renderHeader("environment-data-table", report.hh);

    const completeEnvironmentIndex = [];
    report.g.forEach((group, groupIndex) => {
      group[2].forEach((_row, rowIndex) => completeEnvironmentIndex.push([groupIndex, rowIndex]));
    });
    let visibleEnvironmentIndex = completeEnvironmentIndex;
    let environmentPage = 0;

    const searchInput = document.getElementById("environment-search");
    const pageSizeInput = document.getElementById("environment-page-size");
    const previousButton = document.getElementById("environment-previous");
    const nextButton = document.getElementById("environment-next");
    const pageInfo = document.getElementById("environment-page-info");
    const environmentBody = document.querySelector("#environment-data-table tbody");

    function environmentValues(reference) {
      const group = report.g[reference[0]];
      const compactRow = group[2][reference[1]];
      const metadata = group[0].map(stringValue);
      const offsets = compactRow.length > 5 ? compactRow[5] : group[1];
      const sampleValues = compactRow[3].map(stringValue);
      const changed = new Set(compactRow[4]);
      const values = metadata.concat(compactRow.slice(0, 3).map(stringValue));
      const classes = Array(values.length).fill("");
      for (let sample = 0; sample < report.m; sample += 1) {
        values.push(sample < offsets.length ? offsets[sample] : "");
        values.push(sample < sampleValues.length ? sampleValues[sample] : "");
        classes.push("");
        classes.push(changed.has(sample) ? "changed" : "");
      }
      return {values, classes, metadata};
    }

    function renderEnvironmentPage() {
      const requestedPageSize = Number(pageSizeInput.value);
      const pageSize = requestedPageSize === 0
        ? Math.max(1, visibleEnvironmentIndex.length)
        : requestedPageSize;
      const pageCount = Math.max(1, Math.ceil(visibleEnvironmentIndex.length / pageSize));
      environmentPage = Math.max(0, Math.min(environmentPage, pageCount - 1));
      const start = environmentPage * pageSize;
      const end = Math.min(start + pageSize, visibleEnvironmentIndex.length);
      const fragment = document.createDocumentFragment();
      for (const reference of visibleEnvironmentIndex.slice(start, end)) {
        const rendered = environmentValues(reference);
        const row = appendRow(
          fragment,
          rendered.values,
          rendered.classes,
          rendered.metadata[7] === "exact" ? "" : "inferred",
        );
        row.dataset.reference = rendered.metadata[0];
        row.dataset.process = rendered.metadata[4];
        row.dataset.event = rendered.metadata[5];
        row.dataset.mapping = rendered.metadata[7];
      }
      environmentBody.replaceChildren(fragment);
      pageInfo.textContent = visibleEnvironmentIndex.length
        ? `${start + 1}-${end} of ${visibleEnvironmentIndex.length.toLocaleString()} rows`
        : "0 rows";
      previousButton.disabled = environmentPage === 0;
      nextButton.disabled = environmentPage >= pageCount - 1;
    }

    function applyEnvironmentSearch() {
      const query = searchInput.value.trim().toLocaleLowerCase();
      if (!query) {
        visibleEnvironmentIndex = completeEnvironmentIndex;
      } else {
        visibleEnvironmentIndex = completeEnvironmentIndex.filter(reference => {
          const rendered = environmentValues(reference);
          return rendered.values.some(value => String(value).toLocaleLowerCase().includes(query));
        });
      }
      environmentPage = 0;
      renderEnvironmentPage();
    }

    document.getElementById("environment-search-button").addEventListener("click", applyEnvironmentSearch);
    document.getElementById("environment-clear-button").addEventListener("click", () => {
      searchInput.value = "";
      applyEnvironmentSearch();
    });
    searchInput.addEventListener("keydown", event => {
      if (event.key === "Enter") applyEnvironmentSearch();
    });
    pageSizeInput.addEventListener("change", () => {
      environmentPage = 0;
      renderEnvironmentPage();
    });
    previousButton.addEventListener("click", () => {
      environmentPage -= 1;
      renderEnvironmentPage();
    });
    nextButton.addEventListener("click", () => {
      environmentPage += 1;
      renderEnvironmentPage();
    });
    renderEnvironmentPage();
  } catch (error) {
    const errorBox = document.getElementById("report-error");
    errorBox.hidden = false;
    errorBox.textContent = `The report data could not be displayed: ${error.message}`;
    console.error(error);
  }
})();
</script>
</body>
</html>
"""
        )

    if progress:
        progress(f"Saved {output}")
    return output


def export_xlsx(
    dataset: DecodedDataset,
    path: str | Path,
    progress: Callable[[str], None] | None = None,
) -> Path:
    output = Path(path)
    workbook = Workbook(write_only=True)

    events_sheet = workbook.create_sheet("Event List")
    events_sheet.freeze_panes = "A2"
    _header_row(events_sheet, EVENT_HEADERS)
    for record in dataset.edv.records:
        events_sheet.append(event_export_row(dataset, record))
    events_sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(EVENT_HEADERS))}{len(dataset.edv.records) + 1}"
    )
    _set_widths(
        events_sheet,
        [9, 20, 20, 12, 48, 10, 20, 34, 24, 9, 9, 8, 8, 9, 12, 22, 8]
        + [12] * (len(EVENT_HEADERS) - 17),
    )
    if progress:
        progress("Event list written")

    max_samples = max(
        (record.environment_block_count_start for record in dataset.edv.records), default=0
    )
    environment_headers = [
        "Ref Nr",
        "Start Time",
        "Event Name",
        "Event Description",
        "Process",
        "Event Id",
        "Env Block",
        "Mapping",
        "Variable Name",
        "Description",
        "Unit",
    ]
    for sample in range(1, max_samples + 1):
        environment_headers.extend([f"Offset {sample} (ms)", f"Value {sample}"])

    environment_sheet = workbook.create_sheet("Environment Data")
    environment_sheet.freeze_panes = "A2"
    _header_row(environment_sheet, environment_headers)
    environment_row_count = 1
    for record_index, record in enumerate(dataset.edv.records, 1):
        event_name = (
            record.definition.name
            if record.definition
            else f"Unknown {record.process_id:02X}:{record.event_id:03X}"
        )
        event_description = record.definition.description if record.definition else ""
        block_name = record.environment_block.block_id if record.environment_block else ""
        for environment in environment_rows(record):
            values: list[object] = [
                f"{record.reference_number:05d}",
                _time(record.start_time),
                event_name,
                event_description,
                f"{record.process_id:02X}",
                f"{record.event_id:03X}",
                block_name,
                record.environment_mapping,
                environment.name,
                environment.description,
                environment.unit,
            ]
            cells = [WriteOnlyCell(environment_sheet, value=value) for value in values]
            if record.environment_mapping != "exact":
                for cell in cells:
                    cell.fill = INFERRED_FILL
            for sample in range(max_samples):
                if sample < len(environment.values):
                    offset_cell = WriteOnlyCell(
                        environment_sheet, value=environment.offsets_ms[sample]
                    )
                    value_cell = WriteOnlyCell(
                        environment_sheet, value=environment.values[sample]
                    )
                    if environment.changed[sample]:
                        value_cell.fill = CHANGED_FILL
                    cells.extend([offset_cell, value_cell])
                else:
                    cells.extend(
                        [WriteOnlyCell(environment_sheet, value=""), WriteOnlyCell(environment_sheet, value="")]
                    )
            environment_sheet.append(cells)
            environment_row_count += 1
        if progress and record_index % 50 == 0:
            progress(f"Environment data: {record_index}/{len(dataset.edv.records)} events")
    environment_sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(environment_headers))}{environment_row_count}"
    )
    _set_widths(
        environment_sheet,
        [9, 20, 32, 48, 9, 9, 24, 32, 30, 44, 10] + [14, 14] * max_samples,
    )

    summary_sheet = workbook.create_sheet("Fault Summary")
    _header_row(summary_sheet, ["Count", "Process", "Event Id", "Event Name", "Description"])
    counts = Counter((record.process_id, record.event_id) for record in dataset.edv.records)
    for (process_id, event_id), count in counts.most_common():
        definition = dataset.definitions.events.get((process_id, event_id))
        summary_sheet.append(
            [
                count,
                f"{process_id:02X}",
                f"{event_id:03X}",
                definition.name if definition else "",
                definition.description if definition else "Definition unavailable",
            ]
        )
    _set_widths(summary_sheet, [10, 10, 10, 38, 60])

    repairs_sheet = workbook.create_sheet("Repair Text")
    _header_row(repairs_sheet, ["Event Name", "Event Description", "Repair / Cause / Remedy"])
    used_names: set[str] = set()
    for record in dataset.edv.records:
        if not record.definition or record.definition.name in used_names:
            continue
        used_names.add(record.definition.name)
        details = dataset.texts.details_for(record.definition.name)
        repairs_sheet.append([record.definition.name, record.definition.description, details])
    _set_widths(repairs_sheet, [40, 60, 100])

    info_sheet = workbook.create_sheet("File Information")
    _header_row(info_sheet, ["Field", "Value"])
    for key, value in file_information_rows(dataset):
        info_sheet.append([key, value])
    for warning in dataset.warnings:
        info_sheet.append(["Warning", warning])
    _set_widths(info_sheet, [30, 110])

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    if progress:
        progress(f"Saved {output}")
    return output
